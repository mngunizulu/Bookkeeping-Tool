from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from .db import AppPaths, Database
from .i18n import category_label, kind_label, payment_label, tr


def euro(cents: int | float) -> str:
    return f"{float(cents)/100:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def report_data(db: Database, start: str, end: str):
    rows = db.list_transactions(start, end)
    totals = db.dashboard_totals(start, end)
    profit_net = totals["income_net"] - totals["expense_net"]
    vat_balance = totals["output_vat"] - totals["input_vat"]
    return rows, totals, profit_net, vat_balance


def _language(db: Database) -> str:
    return db.get_language()


def export_csv(db: Database, paths: AppPaths, start: str, end: str) -> Path:
    rows, _, _, _ = report_data(db, start, end)
    lang = _language(db)
    out = paths.exports / f"bookkeeping_{start}_{end}.csv"
    headers = [
        tr("date", lang), tr("type", lang), tr("counterparty", lang), tr("category", lang),
        tr("description", lang), tr("invoice_receipt_no", lang), tr("net_amount", lang),
        tr("vat_rate", lang), tr("vat", lang), tr("gross_amount", lang), tr("payment", lang),
        tr("paid", lang), tr("docs", lang),
    ]
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(headers)
        for r in rows:
            w.writerow([
                r["booking_date"], kind_label(r["kind"], lang), r.get("counterparty_name") or "",
                category_label(r.get("category_name") or "", lang), r["description"],
                r.get("document_number") or "", f"{r['net_amount_cents']/100:.2f}", r["vat_rate"],
                f"{r['vat_amount_cents']/100:.2f}", f"{r['gross_amount_cents']/100:.2f}",
                payment_label(r.get("payment_method") or "", lang), tr("yes", lang) if r["paid"] else tr("no", lang),
                r["document_count"],
            ])
    return out


def export_excel(db: Database, paths: AppPaths, start: str, end: str) -> Path:
    rows, totals, profit_net, vat_balance = report_data(db, start, end)
    lang = _language(db)
    out = paths.exports / f"bookkeeping_{start}_{end}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Buchungen" if lang == "de" else "Transactions"
    headers = [
        tr("date", lang), tr("type", lang), tr("counterparty", lang), tr("category", lang),
        tr("description", lang), tr("invoice_receipt_no", lang), f"{tr('net', lang)} €", f"{tr('vat', lang)} %",
        f"{tr('vat', lang)} €", f"{tr('gross', lang)} €", tr("payment", lang), tr("paid", lang), tr("docs", lang),
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([
            r["booking_date"], kind_label(r["kind"], lang), r.get("counterparty_name") or "",
            category_label(r.get("category_name") or "", lang), r["description"], r.get("document_number") or "",
            r["net_amount_cents"]/100, r["vat_rate"], r["vat_amount_cents"]/100, r["gross_amount_cents"]/100,
            payment_label(r.get("payment_method") or "", lang), bool(r["paid"]), r["document_count"],
        ])
    for col in (7, 9, 10):
        for cells in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cells:
                c.number_format = '#,##0.00 [$€-407]'
    widths = [12, 12, 24, 26, 34, 20, 14, 10, 14, 14, 16, 10, 8]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    s = wb.create_sheet("Übersicht" if lang == "de" else "Summary")
    s.append([tr("period", lang), f"{start} – {end}"])
    s.append([tr("income_net", lang), totals["income_net"]/100])
    s.append([tr("expenses_net", lang), totals["expense_net"]/100])
    s.append([tr("profit_net", lang), profit_net/100])
    s.append([tr("output_vat", lang), totals["output_vat"]/100])
    s.append([tr("input_vat", lang), totals["input_vat"]/100])
    s.append([tr("vat_balance", lang), vat_balance/100])
    s.append([f"{tr('revenue', lang)} ({tr('gross', lang)})", totals["income_gross"]/100])
    s.append([f"{tr('expenses', lang)} ({tr('gross', lang)})", totals["expense_gross"]/100])
    for row in s.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = '#,##0.00 [$€-407]'
    s.column_dimensions['A'].width = 28
    s.column_dimensions['B'].width = 20
    wb.save(out)
    return out


def export_pdf(db: Database, paths: AppPaths, start: str, end: str) -> Path:
    rows, totals, profit_net, vat_balance = report_data(db, start, end)
    company = db.get_company()
    lang = _language(db)
    out = paths.exports / f"bookkeeping_report_{start}_{end}.pdf"
    doc = SimpleDocTemplate(
        str(out), pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm,
        topMargin=12*mm, bottomMargin=12*mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"{company['company_name']} – {tr('bookkeeping_report', lang)}", styles['Title']),
        Paragraph(f"{tr('period', lang)}: {start} – {end}", styles['Normal']),
        Spacer(1, 5*mm),
    ]
    summary = [
        [tr("income_net", lang), euro(totals['income_net']), tr("expenses_net", lang), euro(totals['expense_net']), tr("profit_net", lang), euro(profit_net)],
        [tr("output_vat", lang), euro(totals['output_vat']), tr("input_vat", lang), euro(totals['input_vat']), tr("vat_balance", lang), euro(vat_balance)],
    ]
    st = Table(summary, colWidths=[30*mm, 28*mm]*3)
    st.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.4, colors.grey), ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('BACKGROUND', (0,0), (-1,-1), colors.whitesmoke),
    ]))
    story += [st, Spacer(1, 5*mm)]
    data = [[tr("date",lang), tr("type",lang), tr("counterparty",lang), tr("category",lang), tr("description",lang), tr("net",lang), tr("vat",lang), tr("gross",lang), tr("docs",lang)]]
    for r in rows[:5000]:
        data.append([
            r['booking_date'], kind_label(r['kind'], lang), (r.get('counterparty_name') or '')[:24],
            category_label(r.get('category_name') or '', lang)[:26], r['description'][:34], euro(r['net_amount_cents']),
            euro(r['vat_amount_cents']), euro(r['gross_amount_cents']), str(r['document_count']),
        ])
    table = Table(data, repeatRows=1, colWidths=[21*mm,15*mm,31*mm,34*mm,51*mm,22*mm,20*mm,22*mm,10*mm])
    table.setStyle(TableStyle([
        ('GRID',(0,0),(-1,-1),0.25,colors.lightgrey), ('BACKGROUND',(0,0),(-1,0),colors.grey),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white), ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),7), ('VALIGN',(0,0),(-1,-1),'TOP'),
    ]))
    story.append(table)
    doc.build(story)
    return out
